"""Rendering a draft to DOCX — T043, T043a, T046, T012.

**Rendered once, at generation.** The bytes written here are hashed, and that
hash is what a human approves and what the send path compares against. Nothing
re-renders at send time: a renderer upgrade between approval and send would
otherwise change the attachment under an approval that never saw it.

ATS-safe by construction: one column, standard ``Heading 1`` headings, plain
paragraphs, and **nothing** in headers, footers, text boxes, tables or images —
an extractor reads a document top to bottom, and anything outside that flow is
where parsed CVs lose their contact details.

**T012 — no network at render time.** The output is DOCX written by
``python-docx``, a pure-Python library. There is no PDF engine and no browser in
this path, so nothing here can resolve a remote host; the egress test renders a
document under the socket guard to keep it that way. A PDF step, if one is ever
added, must pass the same test.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from career_scout.documents.generate import BULLET, HEADING, PARAGRAPH, Draft
from career_scout.documents.style_variant import StyleVariant

_TITLES = {
    "cv": None,
    "cover_letter": "Cover letter",
    "motivation_letter": "Motivation letter",
    "worksheet": "Application worksheet",
}

_ALIGN = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT}


def _rule_below(paragraph, weight_pt: float) -> None:
    """A bottom border on the heading paragraph — a line, not a drawing object."""
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(int(weight_pt * 8)))   # eighths of a point
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "444444")
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)  # noqa: SLF001


def _font(run, name: str, size: float, *, bold: bool = False, caps: bool = False,
          small_caps: bool = False) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)  # noqa: SLF001
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.all_caps = caps
    run.font.small_caps = small_caps
    run.font.color.rgb = RGBColor(0, 0, 0)


def render_docx(draft: Draft, variant: StyleVariant, path: Path) -> str:
    """Write ``draft`` to ``path`` and return the SHA-256 of the bytes written."""
    document = Document()
    body = document.styles["Normal"]
    body.font.name = variant.body_font
    body.font.size = Pt(variant.body_size_pt)
    body.paragraph_format.line_spacing = variant.line_spacing
    body.paragraph_format.space_after = Pt(variant.space_after_paragraph_pt)
    # Hyphenation off (T043a): a break inside a surname or a tool name is a
    # parsing hazard bought for a tidier edge.
    settings = document.settings.element
    auto = OxmlElement("w:autoHyphenation")
    auto.set(qn("w:val"), "0")
    settings.append(auto)

    treatment = variant.heading_treatment
    title = _TITLES.get(draft.doc_type)
    if title:
        heading = document.add_paragraph(style="Heading 1")
        _font(heading.add_run(title), variant.heading_font, variant.heading_size_pt, bold=True)

    first_line = True
    for line in draft.lines:
        if line.kind == HEADING and draft.doc_type == "cv" and first_line:
            # The candidate's name: a plain large paragraph, not a heading, so an
            # extractor does not file it as a section title.
            paragraph = document.add_paragraph()
            _font(paragraph.add_run(line.text), variant.heading_font, variant.name_size_pt,
                  bold=True)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.kind == HEADING and line.origin == "template":
            paragraph = document.add_paragraph(style="Heading 1")
            paragraph.paragraph_format.space_before = Pt(variant.section_gap_pt)
            _font(paragraph.add_run(line.text), variant.heading_font, variant.heading_size_pt,
                  bold=treatment.get("bold", True), caps=treatment.get("all_caps", False),
                  small_caps=treatment.get("small_caps", False))
            if treatment.get("rule"):
                _rule_below(paragraph, variant.rule_weight_pt)
            paragraph.alignment = _ALIGN[variant.justification.headings]
        elif line.kind == HEADING:
            # A role or project line: bold body text, left-aligned.
            paragraph = document.add_paragraph()
            _font(paragraph.add_run(line.text), variant.body_font, variant.body_size_pt,
                  bold=True)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.kind == BULLET:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(14)
            paragraph.paragraph_format.first_line_indent = Pt(-10)
            _font(paragraph.add_run(f"{variant.bullet_glyph} {line.text}"), variant.body_font,
                  variant.body_size_pt)
            paragraph.alignment = _ALIGN[variant.justification.bullets]
        else:
            paragraph = document.add_paragraph()
            _font(paragraph.add_run(line.text), variant.body_font, variant.body_size_pt)
            paragraph.alignment = (
                _ALIGN[variant.justification.body_paragraphs] if line.kind == PARAGRAPH
                else WD_ALIGN_PARAGRAPH.LEFT
            )
        first_line = False

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return hashlib.sha256(path.read_bytes()).hexdigest()


def extract_text(path: Path) -> str:
    """What an ATS reads: every paragraph in body order. Used to test the output."""
    document = Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)
